"""Build the Lab-410 Task 2 submission document (docx)."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def title_block():
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Lab-410 — Compiler Construction Lab"); r.bold = True; r.font.size = Pt(16)
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Task 2: Exploring g++ Flags, Underlying Tools & Python's Bytecode")
    r.bold = True; r.font.size = Pt(13)
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.add_run("Fall 2026  |  CLO1").font.size = Pt(11)
    doc.add_paragraph()

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = ""
        run = cell.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(10)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = table.rows[ri].cells[ci]; cell.text = ""
            run = cell.paragraphs[0].add_run(val); run.font.size = Pt(10)
    doc.add_paragraph()

def qa(question, answer):
    p = doc.add_paragraph(); p.add_run(question).bold = True
    doc.add_paragraph(answer)

title_block()
p = doc.add_paragraph(); p.add_run("Name: ").bold = True; p.add_run("Usman hassan")
p = doc.add_paragraph(); p.add_run("Enrollment No: ").bold = True; p.add_run("f2022-003")

doc.add_heading("Part A — g++ pipeline, stage by stage", level=1)
doc.add_paragraph(
    "Using demo.cpp (prints \"Sum: 15\"), each stage of g++'s pipeline was triggered "
    "manually. The four forms of the program observed on disk:")
add_table(
    ["g++ flag", "Stage", "Output file", "Observed"],
    [["-E", "Preprocessing", "demo.i (907,753 B)",
      "Pure text; #include <iostream> fully expanded; own main() intact at the end"],
     ["-S", "Compilation proper", "demo.s (1,726 B)",
      "x86-64 assembly (movl, call, pushq, …)"],
     ["-c", "Assembling", "demo.o (2,224 B)",
      "ELF64 machine code (binary); objdump -d disassembles it back to readable form"],
     ["(none)", "Linking", "demo (16,544 B)",
      "Runnable executable; ./demo prints \"Sum: 15\""]])

doc.add_heading("Part B — the programs g++ actually calls", level=1)
doc.add_paragraph(
    "g++ -v reveals the real tools invoked. cc1plus does the actual compiling "
    "(preprocessing + compilation proper), as assembles, and collect2 wraps ld, "
    "the linker.")
add_table(
    ["Tool", "Role", "Location"],
    [["cc1plus", "Real compiler (stages 1 & 2)",
      "/usr/libexec/gcc/x86_64-linux-gnu/13/cc1plus — via g++ -print-prog-name; NOT on PATH"],
     ["as", "Assembler (stage 3)", "/usr/bin/as — on PATH (which as)"],
     ["collect2", "Wrapper calling ld, the linker (stage 4)",
      "/usr/libexec/gcc/x86_64-linux-gnu/13/collect2 — via g++ -print-prog-name; NOT on PATH"],
     ["ld", "The actual linker", "/usr/bin/ld — on PATH (which ld)"]])
doc.add_paragraph(
    "which cc1plus and which collect2 print nothing: they are GCC-private helpers, "
    "hidden in GCC's libexec directory, while as and ld are public standalone tools.")

doc.add_heading("Part C — Python's bytecode", level=1)
doc.add_paragraph(
    "python3 -m dis demo.py shows CPython compiles the source to bytecode before "
    "executing it — e.g. LOAD_FAST, BINARY_OP, RETURN_VALUE for add() (Python 3.12 "
    "instruction names). Running run_demo.py (which does import demo) produced "
    "__pycache__/demo.cpython-312.pyc (314 B); disassembling that .pyc directly "
    "yields the identical instruction sequence, proving it is the saved compiled "
    "form. The directly-run script is compiled in memory only — no .pyc is written "
    "for it.")

doc.add_heading("Part D — pipeline comparison", level=1)
add_table(
    ["Stage concept", "C++ (g++)", "Python"],
    [["Is there a visible intermediate text form?",
      "Yes — demo.i (preprocessed source) and demo.s (assembly)",
      "No — bytecode is binary/in-memory; dis only displays it"],
     ["Is there a lower-level instruction form?",
      "Yes — assembly (demo.s), then machine code (demo.o)",
      "Yes — bytecode (LOAD_FAST, BINARY_OP, …)"],
     ["Is a separate file saved to disk automatically?",
      "No — intermediates need -E/-S/-c; plain g++ keeps only the executable",
      "For imported modules, yes — __pycache__/*.pyc is written automatically"],
     ["Who executes the final form?", "The CPU, directly",
      "The CPython virtual machine"],
     ["Do you need the original toolchain present to run the final form again?",
      "No — the linked executable runs standalone",
      "Yes — the interpreter/VM is always required"]])

doc.add_heading("Reflection Questions", level=1)
qa("1. In Part A, which g++ flag corresponds to which of the 4 stages "
   "(preprocessing, compilation, assembling, linking)? List all four.",
   "-E → preprocessing; -S → compilation proper (source to assembly); "
   "-c → assembling (assembly to object/machine code); no flag (default "
   "g++ invocation) → linking (object file to executable).")
qa("2. Why do you think cc1plus and collect2 aren't on your PATH, while as "
   "and ld are?",
   "as and ld are general-purpose standalone tools that any toolchain or user "
   "can invoke directly (e.g. assembling hand-written assembly), so they are "
   "installed on PATH. cc1plus and collect2 are GCC's private internal helpers "
   "that only make sense as part of g++'s own pipeline, so GCC keeps them hidden "
   "in its libexec directory and manages them itself.")
qa("3. Based on Part C, is it accurate to say Python has \"no compilation step "
   "at all\"? Why or why not — what actually happens before your Python code runs?",
   "No, it is not accurate. Before Python code runs, CPython compiles the source "
   "to bytecode — an intermediate instruction set visible with the dis module "
   "(e.g. LOAD_FAST, BINARY_OP, RETURN_VALUE) — and the CPython virtual machine "
   "then executes that bytecode. So Python does have a compilation step; it is "
   "just automatic, usually in-memory, and targets a virtual machine instead of "
   "real machine code.")
qa("4. What is the key difference between C++'s object/executable file and "
   "Python's .pyc bytecode file, even though both are described as some kind of "
   "\"compiled\" output?",
   "C++'s .o and executable contain real machine code for the CPU, with needed "
   "library code linked in — the CPU runs them directly. Python's .pyc contains "
   "bytecode for the CPython virtual machine: it is not CPU-executable, nothing "
   "is linked into it, and it still requires the interpreter to run.")
qa("5. Does Python's .pyc file mean Python no longer needs an interpreter to "
   "run the program? Explain.",
   "No. The .pyc holds bytecode, which only the CPython virtual machine "
   "understands — the interpreter is still required to execute it. The .pyc "
   "merely lets Python skip the source-to-bytecode compilation step when an "
   "unchanged module is imported again.")

doc.add_heading("Conclusion", level=1)
doc.add_paragraph(
    "Walking the pipeline by hand shows g++ is a driver coordinating four real "
    "stages — preprocessing, compilation, assembling, linking — carried out by "
    "specialized tools (cc1plus, as, collect2/ld), two of which GCC keeps private. "
    "Python, meanwhile, is not purely interpreted: it compiles every program to "
    "bytecode first and caches that bytecode in .pyc files for imported modules. "
    "The fundamental difference is the compilation target: C++ produces machine "
    "code the CPU runs directly, while Python produces bytecode that only its own "
    "virtual machine can execute.")

out = "/home/hatch/workspace/lab410-task2/Lab410_Task2_Report_f2022-003.docx"
doc.save(out)
print("saved", out)
